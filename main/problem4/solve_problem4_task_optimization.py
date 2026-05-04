from __future__ import annotations

import json
import sys
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.optimize import Bounds, LinearConstraint, milp
from scipy.sparse import lil_matrix


PROJECT_ROOT = Path(__file__).resolve().parents[2]
TRAJECTORY_PATH = PROJECT_ROOT / "outputs" / "问题三" / "problem3_trajectory_10hz_kalman.csv"
TARGET_PATH = PROJECT_ROOT / "数据以及可视化" / "附件4.xlsx"
RESULT_TEMPLATE_PATH = PROJECT_ROOT / "result.xlsx"
OUT_DIR = PROJECT_ROOT / "outputs" / "问题四"

DT = 0.1

SHOOT_DISTANCE_RANGE_M = (5.0, 30.0)
PHOTO_DISTANCE_RANGE_M = (10.0, 40.0)
SHOOT_SPEED_MAX_MPS = 2.0
PHOTO_SPEED_MAX_MPS = 1.5
ACCEL_MAX_MPS2 = 1.5
SHOOT_PREP_TIME_S = 1.5
PHOTO_PREP_TIME_S = 0.5
SHOOT_PREP_POINTS = int(round(SHOOT_PREP_TIME_S / DT)) + 1
PHOTO_PREP_POINTS = int(round(PHOTO_PREP_TIME_S / DT)) + 1
PHOTO_MIN_ANGLE_DIFF_DEG = 60.0

SHOOT_WEIGHT = 0.85
PHOTO_WEIGHT = 1.0
OBJECTIVE_MODE = "最大化期望任务数"

KALMAN_Q = np.diag([0.01, 0.01, 0.001, 0.01, 0.01, 0.001])
KALMAN_R = np.diag([0.1, 0.1])


@dataclass
class CandidateTask:
    target_id: str
    task_type: str
    start_idx: int
    exec_idx: int
    start_time: float
    exec_time: float
    distance_m: float
    speed_mps: float
    accel_mps2: float
    angle_deg: float | None
    weight: float


def load_trajectory() -> pd.DataFrame:
    df = pd.read_csv(TRAJECTORY_PATH)
    required = {"time_s", "x_kalman_m", "y_kalman_m"}
    missing = required - set(df.columns)
    if missing:
        raise ValueError(f"轨迹文件缺少必要列：{sorted(missing)}")
    return df


def load_targets() -> tuple[pd.DataFrame, pd.DataFrame]:
    shoot = pd.read_excel(TARGET_PATH, sheet_name=0).iloc[:, :3].copy()
    photo = pd.read_excel(TARGET_PATH, sheet_name=1).iloc[:, :3].copy()
    shoot.columns = ["target_id", "x_m", "y_m"]
    photo.columns = ["target_id", "x_m", "y_m"]
    return shoot, photo


def kalman_motion_state(x_obs: np.ndarray, y_obs: np.ndarray, dt: float) -> np.ndarray:
    n = len(x_obs)
    if n < 2:
        raise ValueError("轨迹点太少，无法估计运动状态。")

    state = np.zeros((6, n), dtype=float)
    state[0, 0] = float(x_obs[0])
    state[3, 0] = float(y_obs[0])

    f = np.array(
        [
            [1.0, dt, 0.5 * dt**2, 0.0, 0.0, 0.0],
            [0.0, 1.0, dt, 0.0, 0.0, 0.0],
            [0.0, 0.0, 1.0, 0.0, 0.0, 0.0],
            [0.0, 0.0, 0.0, 1.0, dt, 0.5 * dt**2],
            [0.0, 0.0, 0.0, 0.0, 1.0, dt],
            [0.0, 0.0, 0.0, 0.0, 0.0, 1.0],
        ]
    )
    h = np.array(
        [
            [1.0, 0.0, 0.0, 0.0, 0.0, 0.0],
            [0.0, 0.0, 0.0, 1.0, 0.0, 0.0],
        ]
    )

    p = np.eye(6) * 10.0
    eye = np.eye(6)
    for k in range(1, n):
        state_pred = f @ state[:, k - 1]
        p_pred = f @ p @ f.T + KALMAN_Q
        z = np.array([x_obs[k], y_obs[k]], dtype=float)
        gain = p_pred @ h.T @ np.linalg.inv(h @ p_pred @ h.T + KALMAN_R)
        state[:, k] = state_pred + gain @ (z - h @ state_pred)
        p = (eye - gain @ h) @ p_pred

    return state


def build_motion_table(df: pd.DataFrame) -> pd.DataFrame:
    t = df["time_s"].to_numpy(float)
    dt = float(np.median(np.diff(t)))
    if not np.isclose(dt, DT, atol=1e-6):
        print(f"提示：轨迹采样间隔中位数为 {dt:.6f} s，后续计算采用该间隔。")

    x_obs = df["x_kalman_m"].to_numpy(float)
    y_obs = df["y_kalman_m"].to_numpy(float)
    state = kalman_motion_state(x_obs, y_obs, dt)

    return pd.DataFrame(
        {
            "time_s": t,
            "x_m": state[0],
            "y_m": state[3],
            "vx_mps": state[1],
            "vy_mps": state[4],
            "ax_mps2": state[2],
            "ay_mps2": state[5],
            "speed_mps": np.hypot(state[1], state[4]),
            "accel_mps2": np.hypot(state[2], state[5]),
        }
    )


def consecutive_window_ok(mask: np.ndarray, points: int) -> np.ndarray:
    ok = np.zeros(len(mask), dtype=bool)
    for idx in range(points - 1, len(mask)):
        ok[idx] = bool(mask[idx - points + 1 : idx + 1].all())
    return ok


def angle_difference_deg(left: float, right: float) -> float:
    return abs((left - right + 180.0) % 360.0 - 180.0)


def build_shoot_candidates(motion: pd.DataFrame, targets: pd.DataFrame) -> list[CandidateTask]:
    t = motion["time_s"].to_numpy(float)
    x = motion["x_m"].to_numpy(float)
    y = motion["y_m"].to_numpy(float)
    speed = motion["speed_mps"].to_numpy(float)
    accel = motion["accel_mps2"].to_numpy(float)

    candidates: list[CandidateTask] = []
    motion_ok = (speed <= SHOOT_SPEED_MAX_MPS) & (accel <= ACCEL_MAX_MPS2)
    for _, row in targets.iterrows():
        target_id = str(row["target_id"])
        tx = float(row["x_m"])
        ty = float(row["y_m"])
        distance = np.hypot(x - tx, y - ty)
        feasible = (
            motion_ok
            & (distance >= SHOOT_DISTANCE_RANGE_M[0])
            & (distance <= SHOOT_DISTANCE_RANGE_M[1])
        )
        window_ok = consecutive_window_ok(feasible, SHOOT_PREP_POINTS)

        for exec_idx in np.where(window_ok)[0]:
            start_idx = int(exec_idx - SHOOT_PREP_POINTS + 1)
            candidates.append(
                CandidateTask(
                    target_id=target_id,
                    task_type="射击",
                    start_idx=start_idx,
                    exec_idx=int(exec_idx),
                    start_time=float(t[start_idx]),
                    exec_time=float(t[exec_idx]),
                    distance_m=float(distance[exec_idx]),
                    speed_mps=float(speed[exec_idx]),
                    accel_mps2=float(accel[exec_idx]),
                    angle_deg=None,
                    weight=SHOOT_WEIGHT,
                )
            )
    return candidates


def build_photo_candidates(motion: pd.DataFrame, targets: pd.DataFrame) -> list[CandidateTask]:
    t = motion["time_s"].to_numpy(float)
    x = motion["x_m"].to_numpy(float)
    y = motion["y_m"].to_numpy(float)
    speed = motion["speed_mps"].to_numpy(float)
    accel = motion["accel_mps2"].to_numpy(float)

    candidates: list[CandidateTask] = []
    motion_ok = (speed <= PHOTO_SPEED_MAX_MPS) & (accel <= ACCEL_MAX_MPS2)
    for _, row in targets.iterrows():
        target_id = str(row["target_id"])
        tx = float(row["x_m"])
        ty = float(row["y_m"])
        distance = np.hypot(x - tx, y - ty)
        angle = (np.degrees(np.arctan2(ty - y, tx - x)) + 360.0) % 360.0
        feasible = (
            motion_ok
            & (distance >= PHOTO_DISTANCE_RANGE_M[0])
            & (distance <= PHOTO_DISTANCE_RANGE_M[1])
        )
        window_ok = consecutive_window_ok(feasible, PHOTO_PREP_POINTS)

        for exec_idx in np.where(window_ok)[0]:
            start_idx = int(exec_idx - PHOTO_PREP_POINTS + 1)
            candidates.append(
                CandidateTask(
                    target_id=target_id,
                    task_type="拍照",
                    start_idx=start_idx,
                    exec_idx=int(exec_idx),
                    start_time=float(t[start_idx]),
                    exec_time=float(t[exec_idx]),
                    distance_m=float(distance[exec_idx]),
                    speed_mps=float(speed[exec_idx]),
                    accel_mps2=float(accel[exec_idx]),
                    angle_deg=float(angle[exec_idx]),
                    weight=PHOTO_WEIGHT,
                )
            )
    return candidates


def add_constraint(rows: list[list[int]], cols: list[int]) -> None:
    if len(cols) > 1:
        rows.append(cols)


def solve_selection(candidates: list[CandidateTask], n_time: int) -> list[CandidateTask]:
    n = len(candidates)
    if n == 0:
        return []

    rows: list[list[int]] = []

    # 时间片互斥：每个10Hz时刻最多被一个任务窗口占用。
    active_by_time: list[list[int]] = [[] for _ in range(n_time)]
    for col, task in enumerate(candidates):
        for idx in range(task.start_idx, task.exec_idx + 1):
            active_by_time[idx].append(col)
    for cols in active_by_time:
        add_constraint(rows, cols)

    # 每个射击目标最多射击一次。
    shoot_targets = sorted({c.target_id for c in candidates if c.task_type == "射击"})
    for target_id in shoot_targets:
        cols = [i for i, c in enumerate(candidates) if c.task_type == "射击" and c.target_id == target_id]
        add_constraint(rows, cols)

    # 同一拍照目标两次拍照必须有足够的圆周角差。
    photo_targets = sorted({c.target_id for c in candidates if c.task_type == "拍照"})
    for target_id in photo_targets:
        cols = [i for i, c in enumerate(candidates) if c.task_type == "拍照" and c.target_id == target_id]
        for a in range(len(cols)):
            for b in range(a + 1, len(cols)):
                left = candidates[cols[a]]
                right = candidates[cols[b]]
                if left.angle_deg is None or right.angle_deg is None:
                    continue
                if angle_difference_deg(left.angle_deg, right.angle_deg) < PHOTO_MIN_ANGLE_DIFF_DEG:
                    rows.append([cols[a], cols[b]])

    matrix = lil_matrix((len(rows), n), dtype=float)
    lower = np.full(len(rows), -np.inf)
    upper = np.ones(len(rows))
    for row_idx, cols in enumerate(rows):
        matrix[row_idx, cols] = 1.0

    result = milp(
        c=-np.array([c.weight for c in candidates], dtype=float),
        integrality=np.ones(n, dtype=int),
        bounds=Bounds(np.zeros(n), np.ones(n)),
        constraints=LinearConstraint(matrix.tocsr(), lower, upper),
        options={"disp": False, "mip_rel_gap": 0.0},
    )
    if not result.success:
        raise RuntimeError(f"整数规划求解失败：{result.message}")

    selected = [candidates[i] for i, value in enumerate(result.x) if value > 0.5]
    selected.sort(key=lambda item: item.exec_time)
    return selected


def selected_table(selected: list[CandidateTask]) -> pd.DataFrame:
    rows = []
    for idx, task in enumerate(selected, start=1):
        rows.append(
            {
                "序号": idx,
                "目标编号": task.target_id,
                "任务": task.task_type,
                "开始准备时刻(s)": round(task.start_time, 2),
                "任务执行时刻(s)": round(task.exec_time, 2),
                "执行距离(m)": round(task.distance_m, 4),
                "执行速度(m/s)": round(task.speed_mps, 4),
                "执行加速度(m/s2)": round(task.accel_mps2, 4),
                "拍照方向角(deg)": None if task.angle_deg is None else round(task.angle_deg, 2),
                "目标函数权重": task.weight,
            }
        )
    return pd.DataFrame(
        rows,
        columns=[
            "序号",
            "目标编号",
            "任务",
            "开始准备时刻(s)",
            "任务执行时刻(s)",
            "执行距离(m)",
            "执行速度(m/s)",
            "执行加速度(m/s2)",
            "拍照方向角(deg)",
            "目标函数权重",
        ],
    )


def write_result_workbook(table: pd.DataFrame) -> Path:
    wb = load_workbook(RESULT_TEMPLATE_PATH)
    ws = wb.active

    for row in range(2, max(ws.max_row, len(table) + 1) + 1):
        for col in range(1, 6):
            ws.cell(row=row, column=col).value = None

    for row_idx, (_, row) in enumerate(table.iterrows(), start=2):
        ws.cell(row=row_idx, column=1).value = int(row["序号"])
        ws.cell(row=row_idx, column=2).value = row["目标编号"]
        ws.cell(row=row_idx, column=3).value = row["任务"]
        ws.cell(row=row_idx, column=4).value = float(row["开始准备时刻(s)"])
        ws.cell(row=row_idx, column=5).value = float(row["任务执行时刻(s)"])

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    output_copy = OUT_DIR / "result_problem4.xlsx"
    wb.save(output_copy)
    try:
        wb.save(RESULT_TEMPLATE_PATH)
        return RESULT_TEMPLATE_PATH
    except PermissionError:
        return output_copy


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    raw_trajectory = load_trajectory()
    shoot_targets, photo_targets = load_targets()
    motion = build_motion_table(raw_trajectory)
    motion_path = OUT_DIR / "problem4_motion_state_kalman.csv"
    motion.to_csv(motion_path, index=False, encoding="utf-8-sig")

    shoot_candidates = build_shoot_candidates(motion, shoot_targets)
    photo_candidates = build_photo_candidates(motion, photo_targets)
    candidates = shoot_candidates + photo_candidates
    selected = solve_selection(candidates, len(motion))
    table = selected_table(selected)

    selected_path = OUT_DIR / "problem4_selected_tasks.csv"
    table.to_csv(selected_path, index=False, encoding="utf-8-sig")
    workbook_path = write_result_workbook(table)

    n_shoot = int((table["任务"] == "射击").sum()) if not table.empty else 0
    n_photo = int((table["任务"] == "拍照").sum()) if not table.empty else 0
    objective = float(table["目标函数权重"].sum()) if not table.empty else 0.0
    summary = {
        "思路": "轻度卡尔曼滤波估计运动状态 -> 对每个满足连续准备时间的执行时刻生成原子候选窗口 -> 用10Hz时间片互斥约束建立0-1整数规划 -> 用圆周角差约束拍照角度 -> 最大化期望任务数",
        "目标函数": f"{SHOOT_WEIGHT}×射击任务数 + {PHOTO_WEIGHT}×拍照任务数",
        "目标函数模式": OBJECTIVE_MODE,
        "轨迹来源": str(TRAJECTORY_PATH),
        "目标来源": str(TARGET_PATH),
        "运动状态CSV": str(motion_path),
        "射击候选任务数": len(shoot_candidates),
        "拍照候选任务数": len(photo_candidates),
        "候选任务总数": len(candidates),
        "选中射击任务数": n_shoot,
        "选中拍照任务数": n_photo,
        "选中任务总数": int(len(table)),
        "目标函数值": objective,
        "约束": {
            "射击距离范围_m": SHOOT_DISTANCE_RANGE_M,
            "拍照距离范围_m": PHOTO_DISTANCE_RANGE_M,
            "射击最大速度_mps": SHOOT_SPEED_MAX_MPS,
            "拍照最大速度_mps": PHOTO_SPEED_MAX_MPS,
            "最大加速度_mps2": ACCEL_MAX_MPS2,
            "射击准备时间_s": SHOOT_PREP_TIME_S,
            "拍照准备时间_s": PHOTO_PREP_TIME_S,
            "射击准备窗口点数": SHOOT_PREP_POINTS,
            "拍照准备窗口点数": PHOTO_PREP_POINTS,
            "同一拍照目标最小圆周角差_deg": PHOTO_MIN_ANGLE_DIFF_DEG,
        },
        "任务CSV": str(selected_path),
        "结果工作簿": str(workbook_path),
    }
    summary_path = OUT_DIR / "problem4_summary.json"
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

    print(json.dumps(summary, ensure_ascii=False, indent=2))
    if table.empty:
        print("没有选中任务。")
    else:
        cols = ["序号", "目标编号", "任务", "开始准备时刻(s)", "任务执行时刻(s)", "拍照方向角(deg)"]
        print(table[cols].to_string(index=False))


if __name__ == "__main__":
    main()
